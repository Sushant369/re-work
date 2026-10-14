"""OOXML helpers: paragraph text, the PowerPoint style-inheritance chain, bullets, fonts.

Bullet resolution is a port of Docling's ``_get_effective_list_marker`` /
``_is_list_item`` (docling/backend/mspowerpoint_backend.py). The same
inheritance chain is reused to resolve the *effective* font size and bold, which
python-pptx does not do: ``run.font.size`` is ``None`` whenever the size is
inherited from the layout or master, which is what breaks naive font-size rules.
"""

from __future__ import annotations

from typing import Iterator, Optional

from lxml import etree
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.oxml.ns import qn

NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "dgm": "http://schemas.openxmlformats.org/drawingml/2006/diagram",
}
MATH_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
DRAWING_2010_NS = "http://schemas.microsoft.com/office/drawing/2010/main"
MC_NS = "http://schemas.openxmlformats.org/markup-compatibility/2006"
MC_ALTERNATE_CONTENT = f"{{{MC_NS}}}AlternateContent"

SAFE_XML_PARSER = etree.XMLParser(
    resolve_entities=False, load_dtd=False, no_network=True, dtd_validation=False
)

TITLE_TYPES = {
    PP_PLACEHOLDER.TITLE,
    PP_PLACEHOLDER.CENTER_TITLE,
    PP_PLACEHOLDER.VERTICAL_TITLE,
}
FURNITURE_TYPES = {
    PP_PLACEHOLDER.FOOTER,
    PP_PLACEHOLDER.DATE,
    PP_PLACEHOLDER.SLIDE_NUMBER,
}
BODY_STYLE_TYPES = {
    PP_PLACEHOLDER.BODY,
    PP_PLACEHOLDER.OBJECT,
    PP_PLACEHOLDER.SUBTITLE,
    PP_PLACEHOLDER.VERTICAL_BODY,
    PP_PLACEHOLDER.VERTICAL_OBJECT,
}

DEFAULT_FONT_PT = 18.0  # PowerPoint's size when nothing in the chain sets one


# --------------------------------------------------------------------------- text


def iter_paragraph_elements(element: etree._Element) -> Iterator[etree._Element]:
    """Yield runs, breaks, fields and equations in order (Docling port)."""
    for child in element:
        if child.tag == MC_ALTERNATE_CONTENT:
            branch = child.find(f"{{{MC_NS}}}Fallback")
            for choice in child.findall(f"{{{MC_NS}}}Choice"):
                required = choice.get("Requires", "").split()
                if required and all(
                    choice.nsmap.get(prefix) in {MATH_NS, DRAWING_2010_NS}
                    for prefix in required
                ):
                    branch = choice
                    break
            if branch is not None:
                yield from iter_paragraph_elements(branch)
        elif child.tag in {f"{{{DRAWING_2010_NS}}}m", f"{{{MATH_NS}}}oMathPara"}:
            yield from iter_paragraph_elements(child)
        elif child.tag in {qn("a:r"), qn("a:br"), qn("a:fld"), f"{{{MATH_NS}}}oMath"}:
            yield child


def paragraph_runs(p_el: etree._Element) -> list[tuple[Optional[etree._Element], str]]:
    """Return ``(run element or None, text)`` pairs. Equations become plain text.

    Docling converts OMML to LaTeX with its own converter; here the equation's
    text is kept as-is to avoid that dependency.
    """
    out: list[tuple[Optional[etree._Element], str]] = []
    for child in iter_paragraph_elements(p_el):
        if child.tag == f"{{{MATH_NS}}}oMath":
            out.append((None, "".join(child.itertext())))
        elif child.tag == qn("a:br"):
            out.append((None, " "))
        else:
            t = child.find(qn("a:t"))
            out.append((child, (t.text or "") if t is not None else ""))
    return out


def paragraph_text(p_el: etree._Element) -> str:
    return "".join(text for _, text in paragraph_runs(p_el))


def paragraph_level(p_el: etree._Element) -> int:
    pPr = p_el.find("a:pPr", NS)
    if pPr is not None and "lvl" in pPr.attrib:
        try:
            return int(pPr.get("lvl"))
        except ValueError:
            pass
    return 0


def auto_number_start(p_el: etree._Element) -> int:
    bu_auto = p_el.find("a:pPr/a:buAutoNum", NS)
    if bu_auto is not None and "startAt" in bu_auto.attrib:
        try:
            return int(bu_auto.get("startAt"))
        except ValueError:
            pass
    return 1


# ------------------------------------------------------------- inheritance chain


class StyleContext:
    """Deck-wide style sources that are not reachable from a shape."""

    def __init__(self, prs) -> None:
        self.default_text_style = prs._element.find("p:defaultTextStyle", NS)


def _level_node(container: Optional[etree._Element], lvl: int) -> Optional[etree._Element]:
    """Find ``a:lvl{lvl+1}pPr`` in a txBody's lstStyle or directly in a style node."""
    if container is None:
        return None
    if container.tag == qn("p:txBody"):
        container = container.find("a:lstStyle", NS)
        if container is None:
            return None
    return container.find(f"a:lvl{lvl + 1}pPr", NS)


def _layout_placeholder(shape):
    try:
        layout = shape.part.slide.slide_layout
        return layout.placeholders.get(idx=shape.placeholder_format.idx)
    except Exception:
        return None


def _master_placeholder(shape):
    try:
        ph_type = shape.placeholder_format.type
        master = shape.part.slide.slide_layout.slide_master
    except Exception:
        return None
    if ph_type in TITLE_TYPES:
        ph_type = PP_PLACEHOLDER.TITLE
    elif ph_type in BODY_STYLE_TYPES:
        ph_type = PP_PLACEHOLDER.BODY
    try:
        return master.placeholders.get(ph_type)
    except Exception:
        return None


def _master_text_style(shape) -> Optional[etree._Element]:
    """``p:titleStyle`` / ``p:bodyStyle`` / ``p:otherStyle`` of the slide master.

    Docling maps only TITLE to titleStyle and BODY/OBJECT to bodyStyle; PowerPoint
    also uses titleStyle for CENTER_TITLE and bodyStyle for SUBTITLE, so this
    mapping is slightly wider.
    """
    try:
        master = shape.part.slide.slide_layout.slide_master
        ph_type = shape.placeholder_format.type
    except Exception:
        return None
    tx_styles = master._element.find(".//p:txStyles", NS)
    if tx_styles is None:
        return None
    if ph_type in TITLE_TYPES:
        return tx_styles.find("p:titleStyle", NS)
    if ph_type in BODY_STYLE_TYPES:
        return tx_styles.find("p:bodyStyle", NS)
    return tx_styles.find("p:otherStyle", NS)


def ppr_chain(shape, p_el: etree._Element, lvl: int, ctx: StyleContext):
    """Paragraph-property nodes from most to least specific, as PowerPoint resolves them.

    paragraph pPr -> shape lstStyle -> layout placeholder lstStyle ->
    master placeholder lstStyle -> master txStyles (placeholders) /
    presentation defaultTextStyle (free text boxes).
    """
    chain = [
        ("paragraph", p_el.find("a:pPr", NS)),
        ("shape", _level_node(shape._element.find("p:txBody", NS), lvl)),
    ]
    if shape.is_placeholder:
        layout_ph = _layout_placeholder(shape)
        if layout_ph is not None:
            chain.append(("layout", _level_node(layout_ph._element.find("p:txBody", NS), lvl)))
        master_ph = _master_placeholder(shape)
        if master_ph is not None:
            chain.append(
                ("master_placeholder", _level_node(master_ph._element.find("p:txBody", NS), lvl))
            )
        chain.append(("master_style", _level_node(_master_text_style(shape), lvl)))
    else:
        chain.append(("presentation_default", _level_node(ctx.default_text_style, lvl)))
    return [(src, node) for src, node in chain if node is not None]


# ------------------------------------------------------------------------ bullets


def _parse_bullet(pPr: etree._Element) -> tuple[Optional[bool], Optional[str]]:
    """Docling's ``_parse_bullet_from_paragraph_properties`` (is_list, kind)."""
    if pPr.find("a:buNone", NS) is not None:
        return False, "buNone"
    for kind in ("buChar", "buAutoNum", "buBlip"):
        if pPr.find(f"a:{kind}", NS) is not None:
            return True, kind
    return None, None


def resolve_list_marker(shape, paragraph, ctx: StyleContext) -> tuple[bool, str]:
    """Return ``(is_list, "Bullet" | "Numbered" | "None")`` (Docling ``_is_list_item``).

    The first node in the inheritance chain with an explicit marker decides; with
    none, an indented paragraph (lvl > 0) is still treated as a sub-list item.
    """
    p = paragraph._element
    lvl = paragraph_level(p)
    for _, node in ppr_chain(shape, p, lvl, ctx):
        is_list, kind = _parse_bullet(node)
        if is_list is True:
            return True, {"buChar": "Bullet", "buAutoNum": "Numbered"}.get(kind, "None")
        if is_list is False:
            return False, "None"
    return lvl > 0, "None"


# -------------------------------------------------------------------------- fonts


def _resolve_rpr_attr(r_el, chain, attr: str) -> Optional[str]:
    if r_el is not None:
        rPr = r_el.find("a:rPr", NS)
        if rPr is not None and rPr.get(attr) is not None:
            return rPr.get(attr)
    for _, node in chain:
        d = node.find("a:defRPr", NS)
        if d is not None and d.get(attr) is not None:
            return d.get(attr)
    return None


def autofit_scale(shape) -> float:
    """``normAutofit/@fontScale``: PowerPoint shrinks text to fit and records it here."""
    na = shape._element.find("p:txBody/a:bodyPr/a:normAutofit", NS)
    if na is not None and na.get("fontScale"):
        try:
            return int(na.get("fontScale")) / 100000
        except ValueError:
            pass
    return 1.0


def paragraph_style(shape, p_el, ctx: StyleContext) -> tuple[float, bool, bool]:
    """Effective ``(size_pt, bold, caps)`` of a paragraph.

    Size is the one covering most characters; bold/caps require all visible runs.
    """
    chain = ppr_chain(shape, p_el, paragraph_level(p_el), ctx)
    scale = autofit_scale(shape)
    size_chars: dict[float, int] = {}
    total = bold_chars = caps_chars = 0
    for r_el, text in paragraph_runs(p_el):
        n = len(text.strip())
        if not n or r_el is None:
            continue
        sz = _resolve_rpr_attr(r_el, chain, "sz")
        size = round((int(sz) / 100 if sz else DEFAULT_FONT_PT) * scale, 1)
        size_chars[size] = size_chars.get(size, 0) + n
        if _resolve_rpr_attr(r_el, chain, "b") in ("1", "true"):
            bold_chars += n
        if _resolve_rpr_attr(r_el, chain, "cap") == "all" or text.isupper():
            caps_chars += n
        total += n
    if not total:
        return DEFAULT_FONT_PT * scale, False, False
    size = max(size_chars, key=size_chars.get)
    return size, bold_chars == total, caps_chars == total
