"""Stage 1-2: walk the deck and emit raw elements (port of Docling's ``_walk_linear``).

No heading decisions are made here except Docling's one structural rule: text
in a TITLE / CENTER_TITLE placeholder is the slide title. Everything else is a
paragraph or list item carrying its effective style for the next stage.
"""

from __future__ import annotations

import logging
from dataclasses import dataclass
from pathlib import Path
from typing import Callable, Optional

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.shapes.autoshape import Shape
from pptx.shapes.group import GroupShape
from pptx.shapes.picture import Picture

from .models import Deck, Element
from .ooxml import (
    FURNITURE_TYPES,
    TITLE_TYPES,
    StyleContext,
    auto_number_start,
    autofit_scale,
    paragraph_level,
    paragraph_style,
    paragraph_text,
    resolve_list_marker,
)
from .ordering import Transform, iter_shapes_by_position
from .tables import chart_to_html, is_smartart, smartart_texts, table_to_html

_log = logging.getLogger(__name__)

ImageDescriber = Callable[[Path, Element], str]


@dataclass
class _OpenList:
    level: int
    counter: int = 0


def _is_hidden(shape) -> bool:
    hidden = shape._element.xpath("./*/p:cNvPr/@hidden")
    return bool(hidden) and hidden[0] in ("1", "true")


def _alt_text(shape) -> str:
    descr = shape._element.xpath("./*/p:cNvPr/@descr")
    return descr[0].strip() if descr else ""


class _Walker:
    def __init__(self, prs, source: str, image_dir: Optional[Path], describe_image: Optional[ImageDescriber]):
        self.prs = prs
        self.ctx = StyleContext(prs)
        self.slide_w = int(prs.slide_width)
        self.slide_h = int(prs.slide_height)
        self.image_dir = image_dir
        self.describe_image = describe_image
        self.deck = Deck(source, self.slide_w, self.slide_h, len(prs.slides))

    def add(self, el: Element) -> Element:
        el.order = el.uid = len(self.deck.elements)
        self.deck.elements.append(el)
        return el

    def _base(self, shape, slide_no: int, xf: Transform, layer: str, kind: str, **kw) -> Element:
        ph = None
        if getattr(shape, "is_placeholder", False):
            try:
                ph = shape.placeholder_format.type.name
            except Exception:
                ph = None
        return Element(
            slide=slide_no,
            kind=kind,
            bbox=xf.bbox(shape, self.slide_w, self.slide_h),
            layer=layer,
            shape_id=shape.shape_id,
            shape_name=shape.name,
            shape_key=f"{slide_no}:{shape.shape_id}",
            placeholder=ph,
            font_scale=autofit_scale(shape) if getattr(shape, "has_text_frame", False) else 1.0,
            **kw,
        )

    # ------------------------------------------------------------------ walk

    def walk(self) -> Deck:
        for slide_no, slide in enumerate(self.prs.slides, start=1):
            slide_layer = "invisible" if slide._element.get("show") in ("0", "false") else "body"
            for shape in iter_shapes_by_position(slide.shapes):
                self.shape(shape, slide_no, Transform(), slide_layer)

            if slide.has_notes_slide:
                frame = slide.notes_slide.notes_text_frame
                notes = frame.text.strip() if frame is not None else ""
                if notes:
                    self.add(Element(slide=slide_no, kind="notes", text=notes, layer="notes"))
        return self.deck

    def shape(self, shape, slide_no: int, xf: Transform, layer: str) -> None:
        if _is_hidden(shape):
            layer = "invisible"
        if isinstance(shape, GroupShape):
            inner = xf.into_group(shape)
            for child in iter_shapes_by_position(shape.shapes):
                self.shape(child, slide_no, inner, layer)
            return
        try:
            if shape.has_table:
                self.table(shape, slide_no, xf, layer)
            if shape.has_chart:
                self.chart(shape, slide_no, xf, layer)
            if is_smartart(shape):
                self.smartart(shape, slide_no, xf, layer)
            if isinstance(shape, Picture):
                self.picture(shape, slide_no, xf, layer)
            if isinstance(shape, Shape) and shape.has_text_frame and shape.text_frame.text.strip():
                self.text(shape, slide_no, xf, layer)
        except Exception as exc:  # one bad shape must not lose the deck
            _log.warning("slide %d: skipping shape %r: %s", slide_no, shape.name, exc)

    # -------------------------------------------------------------- handlers

    def table(self, shape, slide_no, xf, layer):
        html = table_to_html(shape)
        if html:
            self.add(self._base(shape, slide_no, xf, layer, "table", html=html))

    def chart(self, shape, slide_no, xf, layer):
        title, chart_type, html = chart_to_html(shape)
        self.add(
            self._base(shape, slide_no, xf, layer, "chart", text=title or "", html=html, chart_type=chart_type)
        )

    def smartart(self, shape, slide_no, xf, layer):
        texts = smartart_texts(shape)
        for i, text in enumerate(texts):
            self.add(
                self._base(
                    shape, slide_no, xf, layer, "list_item", text=text,
                    para_index=i, para_count=len(texts), reason="smartart",
                )
            )

    def picture(self, shape, slide_no, xf, layer):
        el = self._base(shape, slide_no, xf, layer, "picture", alt_text=_alt_text(shape))
        try:
            image = shape.image
            if self.image_dir is not None:
                self.image_dir.mkdir(parents=True, exist_ok=True)
                path = self.image_dir / f"slide{slide_no:03d}_shape{shape.shape_id}.{image.ext}"
                path.write_bytes(image.blob)
                el.image_path = str(path)
        except Exception as exc:  # empty picture placeholder, broken rel, ...
            _log.debug("slide %d: picture %r has no image: %s", slide_no, shape.name, exc)
        self.add(el)
        if self.describe_image and el.image_path:
            try:
                el.description = self.describe_image(Path(el.image_path), el) or ""
            except Exception as exc:
                _log.warning("image description failed for %s: %s", el.image_path, exc)

    def text(self, shape, slide_no, xf, layer):
        ph_type = shape.placeholder_format.type if shape.is_placeholder else None
        if ph_type in FURNITURE_TYPES and layer == "body":
            layer = "furniture"
        paragraphs = list(shape.text_frame.paragraphs)

        if ph_type in TITLE_TYPES:
            # Docling emits one TITLE item per paragraph; a wrapped title is one heading.
            texts = [t for t in (paragraph_text(p._element).strip() for p in paragraphs) if t]
            size, bold, caps = paragraph_style(shape, paragraphs[0]._element, self.ctx)
            self.add(
                self._base(
                    shape, slide_no, xf, layer, "title", text=" ".join(texts), level=1,
                    font_size=size, bold=bold, caps=caps, reason="placeholder:title",
                )
            )
            return

        open_lists: list[_OpenList] = []
        emitted: list[Element] = []
        for paragraph in paragraphs:
            p = paragraph._element
            text = paragraph_text(p).strip()
            is_list, bullet = resolve_list_marker(shape, paragraph, self.ctx)
            if not text:
                if not is_list:
                    open_lists.clear()
                continue
            size, bold, caps = paragraph_style(shape, p, self.ctx)
            el = self._base(
                shape, slide_no, xf, layer, "paragraph", text=text,
                font_size=size, bold=bold, caps=caps, para_index=len(emitted),
            )
            if is_list:
                # Docling's nesting: a deeper lvl opens a sub-list, a shallower one
                # closes lists down to (but never past) the outermost.
                level = paragraph_level(p)
                while len(open_lists) > 1 and open_lists[-1].level > level:
                    open_lists.pop()
                if not open_lists or level > open_lists[-1].level:
                    open_lists.append(_OpenList(level))
                current = open_lists[-1]
                el.kind = "list_item"
                el.level = len(open_lists) - 1
                if bullet == "Numbered":
                    if current.counter == 0:
                        current.counter = auto_number_start(p) - 1
                    current.counter += 1
                    el.enumerated, el.enum_index = True, current.counter
            else:
                open_lists.clear()
            emitted.append(self.add(el))

        for el in emitted:
            el.para_count = len(emitted)
        if ph_type == PP_PLACEHOLDER.SUBTITLE:
            for el in emitted:
                el.reason = "placeholder:subtitle"


def extract(
    path: str | Path,
    image_dir: Optional[Path] = None,
    describe_image: Optional[ImageDescriber] = None,
) -> Deck:
    prs = Presentation(str(path))
    return _Walker(prs, str(path), image_dir, describe_image).walk()
