"""Regression tests: each test builds a small deck that once broke the pipeline."""

import io
import sys
from pathlib import Path

import pytest
from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

sys.path.insert(0, str(Path(__file__).parent))
from make_sample_deck import main as make_sample, no_placeholder_deck, textbox  # noqa: E402

from pptx_pipeline.extract import extract
from pptx_pipeline.headings import classify, is_kpi
from pptx_pipeline.reading_order import refine_reading_order
from pptx_pipeline.sections import assign_sections, strip_continuation, title_key


def _deck():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    return prs


def _run(prs, tmp_path, name="deck"):
    path = tmp_path / f"{name}.pptx"
    prs.save(path)
    deck = extract(path)
    classify(deck)
    refine_reading_order(deck)
    return deck, assign_sections(deck)


def _texts(deck, slide, kind=None):
    return [
        e.text for e in deck.slide_elements(slide)
        if e.layer == "body" and (kind is None or e.kind == kind)
    ]


# ------------------------------------------------------------------ unit rules


@pytest.mark.parametrize("text", ["$4.2M", "35%", "+12% YoY", "500+ clients", "3x faster", "4.2 Cr", "1,200", "12"])
def test_kpis(text):
    assert is_kpi(text)


@pytest.mark.parametrize("text", ["2025 Strategy", "3 Key Risks", "100 Day Plan", "FY26 plan", "Q3 2025 Results", "Top 5 priorities"])
def test_number_led_headings_are_not_kpis(text):
    assert not is_kpi(text)


@pytest.mark.parametrize("a,b", [
    ("Risks", "Risks (cont.)"), ("Risks", "Risks – continued"), ("Risks", "Risks contd"),
    ("Risks", "Risks 1/2"), ("Risks", "Risks (2/3)"), ("Risks", "Risks - 2 of 3"),
    ("Risks", "Risks, part 2"), ("Risks", "RISKS"),
])
def test_continuation_titles(a, b):
    assert title_key(a) == title_key(b)


@pytest.mark.parametrize("a,b", [
    ("Plan 2025/26", "Plan 2026/27"), ("Wave (1)", "Wave (2)"), ("Phase 1", "Phase 2"),
])
def test_distinct_titles(a, b):
    assert title_key(a) != title_key(b)


def test_strip_keeps_words_ending_in_cont():
    assert strip_continuation("Discontinued products") == "Discontinued products"


# --------------------------------------------------------------- sample decks


def test_sample_deck(tmp_path):
    make_sample(tmp_path / "sample.pptx")
    deck = extract(tmp_path / "sample.pptx")
    classify(deck)
    refine_reading_order(deck)
    sections = assign_sections(deck)

    assert _texts(deck, 2, "kpi") == ["$4.2M", "35%", "+12% YoY"]
    assert "Key drivers" in _texts(deck, 2, "heading")
    # two-column slide: each heading directly followed by its body
    assert _texts(deck, 3)[1:5] == [
        "Key drivers", "Supplier concentration in one region.",
        "Mitigations", "Dual-source critical parts by Q1.",
    ]
    assert _texts(deck, 5, "title") == ["Regional performance"]  # text-box title
    assert [(s.title, s.slides) for s in sections][-2:] == [
        ("Regional performance", [5, 6, 7, 8]), ("Appendix", [9]),
    ]
    assert deck.slide_elements(7)[0].continued


def test_no_placeholder_deck(tmp_path):
    no_placeholder_deck(tmp_path / "nt.pptx")
    deck = extract(tmp_path / "nt.pptx")
    classify(deck)
    sections = assign_sections(deck)
    assert [(s.title, s.slides) for s in sections] == [("Market overview", [1, 2]), ("Outlook", [3])]
    assert _texts(deck, 2, "heading") == ["Highlights"]  # not promoted to a title


# ------------------------------------------------------------------ edge cases


def test_sentence_titles_without_placeholders(tmp_path):
    prs = _deck()
    for title, body in [
        ("Revenue grew 12% driven by enterprise renewals in all regions this quarter.", "Detail A"),
        ("Margins improved as APAC pricing changes took effect.", "Detail B"),
    ]:
        s = prs.slides.add_slide(prs.slide_layouts[6])
        textbox(s, .8, .3, 11, .8, [(title, 24, False)])
        textbox(s, .8, 1.5, 11, 2, [(body, 14, False)])
    deck, sections = _run(prs, tmp_path)
    assert [s.slides for s in sections] == [[1], [2]]
    assert all(s.title and s.title.endswith(".") for s in sections)


def test_heading_below_kpi_row_is_not_a_label(tmp_path):
    prs = _deck()
    for i in range(2):
        s = prs.slides.add_slide(prs.slide_layouts[5])
        s.shapes.title.text = f"Results {i}"
        textbox(s, .8, 1.6, 3, 1, [("$4.2M", 40, True)])
        textbox(s, .8, 3.0, 11, .5, [("Key drivers", 20, True)])  # 0.4in below the KPI
        textbox(s, .8, 3.6, 11, 1, [("Renewals closed early.", 14, False)])
    deck, _ = _run(prs, tmp_path)
    assert _texts(deck, 1, "heading") == _texts(deck, 2, "heading") == ["Key drivers"]


def test_kpi_label_directly_below(tmp_path):
    prs = _deck()
    for i in range(2):
        s = prs.slides.add_slide(prs.slide_layouts[5])
        s.shapes.title.text = f"Results {i}"
        textbox(s, .8, 1.6, 3, 1, [("$4.2M", 40, True)])
        textbox(s, .8, 2.65, 3, .4, [("Revenue", 14, True)])  # tight label, bold
        textbox(s, .8, 3.6, 11, 1, [("Renewals closed early in all regions.", 14, False)])
    deck, _ = _run(prs, tmp_path)
    label = next(e for e in deck.slide_elements(1) if e.text == "Revenue")
    assert label.kind == "paragraph" and label.reason == "kpi_label"


def test_autofit_keeps_heading_style(tmp_path):
    prs = _deck()
    for i, scale in enumerate([None, None, "85000"]):
        s = prs.slides.add_slide(prs.slide_layouts[5])
        s.shapes.title.text = f"Topic {i}"
        tf = textbox(s, .8, 1.6, 11, 3, [("Key drivers", 20, True), ("Body text here.", 14, False)])
        if scale:
            etree.SubElement(tf._txBody.bodyPr, "{http://schemas.openxmlformats.org/drawingml/2006/main}normAutofit").set("fontScale", scale)
    deck, _ = _run(prs, tmp_path)
    assert [_texts(deck, n, "heading") for n in (1, 2, 3)] == [["Key drivers"]] * 3


def test_background_picture_keeps_column_order(tmp_path):
    prs = _deck()
    buf = io.BytesIO()
    Image.new("RGB", (40, 20), "white").save(buf, "PNG")
    for i in range(2):
        s = prs.slides.add_slide(prs.slide_layouts[5])
        buf.seek(0)
        s.shapes.add_picture(buf, 0, 0, prs.slide_width, prs.slide_height)
        s.shapes.title.text = f"Two columns {i}"
        textbox(s, .8, 1.6, 5.5, .5, [("Left head", 20, True)])
        textbox(s, 7, 1.6, 5.5, .5, [("Right head", 20, True)])
        textbox(s, .8, 2.2, 5.5, 2, [("Left body.", 14, False)])
        textbox(s, 7, 2.2, 5.5, 2, [("Right body.", 14, False)])
    deck, _ = _run(prs, tmp_path)
    assert [t for t in _texts(deck, 1) if t][1:] == ["Left head", "Left body.", "Right head", "Right body."]


def test_fiscal_year_titles_are_separate_sections(tmp_path):
    prs = _deck()
    for title in ["Plan 2025/26", "Plan 2026/27", "Wave (1)", "Wave (2)"]:
        s = prs.slides.add_slide(prs.slide_layouts[5])
        s.shapes.title.text = title
        textbox(s, .8, 1.6, 11, 1, [("Details.", 14, False)])
    _, sections = _run(prs, tmp_path)
    assert len(sections) == 4
